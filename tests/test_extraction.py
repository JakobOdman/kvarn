"""Extraction with the fake client: rows from several documents, errors per document, the cache."""
from io import BytesIO

import pytest
from openpyxl import load_workbook

from api import db, llm
from api.extract import pages_as_text
from tests.conftest import TEMPLATE, add_read_document


class ByCompany(llm.FakeClient):
    """Answers with one row per company name found in the prompt."""

    def send(self, prompt, schema, model):
        self.calls += 1
        rows = [{"bolag": name, "varde": value, "sida": 1}
                for name, value in [("Alfa AB", "100"), ("Beta AB", "250")] if name in prompt]
        return {"innehav": rows}, 1000, 50


@pytest.fixture
def by_company(monkeypatch):
    monkeypatch.setattr(llm, "FakeClient", ByCompany)


def run(client, folder, **request):
    template = client.put("/templates", json=TEMPLATE).json()["id"]
    created = client.post("/extractions", json={"template_id": template, "folder_id": folder, **request})
    assert created.status_code == 200, created.text
    return client.get(f"/extractions/{created.json()['extraction_id']}").json()


def test_rows_from_every_document_end_up_in_one_table(anna, by_company):
    folder = anna.post("/folders", json={"name": "Fonder"}).json()["id"]
    add_read_document(folder, "alfa.pdf", "Alfa AB 100")
    add_read_document(folder, "beta.pdf", "Beta AB 250")
    extraction = run(anna, folder)

    assert extraction["status"] == "done"
    assert sorted((r["dokument"], r["bolag"]) for r in extraction["tables"]["innehav"]) == [
        ("alfa.pdf", "Alfa AB"), ("beta.pdf", "Beta AB")]
    assert [(d["status"], d["tokens_in"], d["tokens_out"]) for d in extraction["documents"]] == [("done", 1000, 50)] * 2


def test_one_bad_document_does_not_stop_the_others(anna, by_company):
    folder = anna.post("/folders", json={"name": "Fonder"}).json()["id"]
    good = add_read_document(folder, "alfa.pdf", "Alfa AB 100")
    unread = add_read_document(folder, "oläst.pdf")
    db.update_document(unread, "error", error="Trasig fil")
    empty = add_read_document(folder, "tom.pdf")
    db.update_document(empty, "done", result={"ok": True, "page_count": 0, "pages": []})
    extraction = run(anna, folder, job_ids=[unread, good, empty])

    assert [(d["name"], d["status"], d["error"]) for d in extraction["documents"]] == [
        ("oläst.pdf", "error", "Dokumentet är inte läst."),
        ("alfa.pdf", "done", None),
        ("tom.pdf", "error", "Dokumentet har ingen text."),
    ]
    assert [r["bolag"] for r in extraction["tables"]["innehav"]] == ["Alfa AB"]


def test_a_folder_runs_only_its_read_documents(anna, by_company):
    folder = anna.post("/folders", json={"name": "Fonder"}).json()["id"]
    add_read_document(folder, "alfa.pdf", "Alfa AB 100")
    db.update_document(add_read_document(folder, "läses.pdf"), "running")
    assert [d["name"] for d in run(anna, folder)["documents"]] == ["alfa.pdf"]


def test_paid_without_the_server_flag_uses_the_fake_client(anna):
    """paid=True from the UI is not enough: without READ_DOCUMENT_PAID=1 no real client is made
    (conftest makes a real client fail the test)."""
    folder = anna.post("/folders", json={"name": "Fonder"}).json()["id"]
    add_read_document(folder)
    assert run(anna, folder, paid=True)["documents"][0]["status"] == "done"


def test_xlsx_has_document_fields_and_page(anna, by_company):
    folder = anna.post("/folders", json={"name": "Fonder"}).json()["id"]
    add_read_document(folder, "alfa.pdf", "Alfa AB 100")
    extraction = run(anna, folder)
    sheet = load_workbook(BytesIO(anna.get(f"/extractions/{extraction['id']}/xlsx").content))["innehav"]
    assert [list(r) for r in sheet.iter_rows(values_only=True)] == [
        ["dokument", "bolag", "varde", "sida"], ["alfa.pdf", "Alfa AB", "100", 1]]


def test_pages_as_text():
    result = {"pages": [{"page_no": 1, "text": "Första", "reader": "pdf_text"},
                        {"page_no": 2, "text": "Andra", "reader": "claude_vision"}]}
    assert pages_as_text(result) == "=== Sida 1 ===\nFörsta\n\n=== Sida 2 ===\nAndra"


# --- The cache ---

SCHEMA = {"type": "object"}


def test_a_cached_answer_needs_no_client():
    db.Cache("anna").put(llm.cache_key("p", SCHEMA, llm.MODEL),
                         {"model": llm.MODEL, "tokens_in": 7, "tokens_out": 3, "answer": {"ok": 1}})
    assert llm.ask("p", SCHEMA, db.Cache("anna")) == ({"ok": 1}, {"tokens_in": 7, "tokens_out": 3, "cached": True})
    with pytest.raises(llm.NotCached):
        llm.ask("p", SCHEMA, db.Cache("bertil"))  # Anna's answer is not Bertil's


class RealLooking:
    """Not the fake client: llm.ask treats it as a paid call."""

    def __init__(self):
        self.calls = 0

    def send(self, prompt, schema, model):
        self.calls += 1
        return {"bolag": "Alfa AB"}, 900, 40


def test_no_cache_and_no_paid_call_is_not_cached():
    client = RealLooking()
    with pytest.raises(llm.NotCached):
        llm.ask("p", SCHEMA, db.Cache("anna"))
    with pytest.raises(llm.NotCached):
        llm.ask("p", SCHEMA, db.Cache("anna"), client=client, paid=True)  # paid, but the server is not in paid mode
    assert client.calls == 0


def test_a_paid_answer_is_saved_and_used_again(monkeypatch):
    monkeypatch.setenv("READ_DOCUMENT_PAID", "1")
    client = RealLooking()
    first = llm.ask("p", SCHEMA, db.Cache("anna"), db.Meter("anna"), client=client, paid=True)
    again = llm.ask("p", SCHEMA, db.Cache("anna"), db.Meter("anna"), client=client, paid=True)
    assert client.calls == 1
    assert first == ({"bolag": "Alfa AB"}, {"tokens_in": 900, "tokens_out": 40, "cached": False})
    assert again == ({"bolag": "Alfa AB"}, {"tokens_in": 900, "tokens_out": 40, "cached": True})


def test_fake_answers_are_never_cached():
    llm.ask("p", SCHEMA, db.Cache("anna"), client=llm.FakeClient({"x": 1}))
    assert db.Cache("anna").get(llm.cache_key("p", SCHEMA, llm.MODEL)) is None
